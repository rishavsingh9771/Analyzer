import io
import re
from typing import Dict, Any, List, Optional
from pypdf import PdfReader
import docx

def extract_text_from_pdf(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text and structural metadata from PDF file bytes."""
    try:
        pdf_stream = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_stream)
        num_pages = len(reader.pages)
        if num_pages == 0:
            raise ValueError("The uploaded PDF has no pages.")
        
        full_text = []
        has_tables_or_multi_column = False

        for page in reader.pages:
            text = page.extract_text() or ""
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            if len(lines) > 20:
                short_lines = sum(1 for l in lines if len(l) < 30)
                if short_lines / len(lines) > 0.6:
                    has_tables_or_multi_column = True
            full_text.append(text)
            
        combined_text = "\n".join(full_text).strip()
        if not combined_text:
            raise ValueError("Could not extract readable text. The document may be an image-only scan.")
            
        return {
            "text": combined_text,
            "page_count": num_pages,
            "has_multi_column_risk": has_tables_or_multi_column,
            "file_type": "pdf"
        }
    except Exception as e:
        raise ValueError(f"Failed to parse PDF document: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> Dict[str, Any]:
    """Extracts text and structural metadata from DOCX file bytes."""
    try:
        docx_stream = io.BytesIO(file_bytes)
        doc = docx.Document(docx_stream)
        
        paragraphs_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        has_tables = len(doc.tables) > 0
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)
                    
        combined = "\n".join(paragraphs_text + table_text).strip()
        if not combined:
            raise ValueError("Could not extract readable text from DOCX.")
            
        return {
            "text": combined,
            "page_count": max(1, len(combined) // 2500),
            "has_multi_column_risk": has_tables,
            "file_type": "docx"
        }
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    """Extracts email, phone, linkedin, github from text."""
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    phone_pattern = r'(\+?\d{1,3}[-.\s]?)?(\(?\d{3}\)?[-.\s]?)?\d{3}[-.\s]?\d{4}'
    linkedin_pattern = r'(?:linkedin\.com\/(?:in\/)?([a-zA-Z0-9-_]+)|(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[^\s]+)'
    github_pattern = r'(?:github\.com\/([a-zA-Z0-9-_]+)|(?:https?:\/\/)?(?:www\.)?github\.com\/[^\s]+)'

    email_match = re.search(email_pattern, text)
    phone_match = re.search(phone_pattern, text)
    linkedin_match = re.search(linkedin_pattern, text, re.IGNORECASE)
    github_match = re.search(github_pattern, text, re.IGNORECASE)

    lines = [l.strip() for l in text.splitlines() if l.strip() and not re.search(email_pattern, l)]
    probable_name = lines[0] if lines and len(lines[0].split()) <= 4 else "Candidate"

    return {
        "name": probable_name,
        "email": email_match.group(0) if email_match else None,
        "phone": phone_match.group(0) if phone_match else None,
        "linkedin": linkedin_match.group(0) if linkedin_match else None,
        "github": github_match.group(0) if github_match else None,
    }

def detect_sections(text: str) -> Dict[str, str]:
    """Segment resume into conventional sections."""
    section_patterns = {
        "summary": r'(?:professional\s+summary|summary|profile|about\s+me|objective)',
        "skills": r'(?:technical\s+skills|core\s+competencies|skills\s*&?\s*tools|skills|technologies)',
        "experience": r'(?:work\s+experience|professional\s+experience|experience|employment\s+history)',
        "projects": r'(?:academic\s+projects|personal\s+projects|key\s+projects|projects)',
        "education": r'(?:education|academic\s+background|qualifications|academic\s+history)',
        "certifications": r'(?:certifications|certificates|licenses|courses)'
    }

    found_sections: Dict[str, str] = {}
    lines = text.splitlines()
    current_section = "intro"
    buffer: List[str] = []

    for line in lines:
        cleaned_line = line.strip().lower()
        matched_section = None
        for sec_name, pattern in section_patterns.items():
            if re.match(rf'^{pattern}[:\s]*$', cleaned_line, re.IGNORECASE) or (len(cleaned_line) < 35 and re.search(rf'\b{pattern}\b', cleaned_line, re.IGNORECASE)):
                matched_section = sec_name
                break
        
        if matched_section:
            if current_section:
                found_sections[current_section] = "\n".join(buffer).strip()
            current_section = matched_section
            buffer = []
        else:
            buffer.append(line)

    if current_section and buffer:
        found_sections[current_section] = "\n".join(buffer).strip()

    return found_sections
